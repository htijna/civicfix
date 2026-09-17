import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def md_to_docx(md_filepath, docx_filepath):
    doc = Document()
    
    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.bold = True
        elif line.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(line[3:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.bold = True
        elif line.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line[4:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
        elif line.startswith('- '):
            p = doc.add_paragraph(line[2:], style='List Bullet')
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
        else:
            p = doc.add_paragraph()
            # Basic bold formatting
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    # Basic italic formatting
                    it_parts = re.split(r'(\*.*?\*)', part)
                    for it_part in it_parts:
                        if it_part.startswith('*') and it_part.endswith('*'):
                            run = p.add_run(it_part[1:-1])
                            run.italic = True
                        else:
                            run = p.add_run(it_part)
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(12)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
                
    doc.save(docx_filepath)

md_to_docx('project_report.md', 'project_report.docx')
